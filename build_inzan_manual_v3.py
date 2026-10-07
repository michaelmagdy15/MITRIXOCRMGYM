"""Inzan handbook: native Word text with an editorial page plan.

Brand reference: https://inzanathletics.com/ inspected 7 October 2026.
Website photography and wordmark are retained in docs/inzan_manual_assets.
"""
from pathlib import Path
from copy import deepcopy
import re
from PIL import Image, ImageDraw, ImageFont, ImageOps, ImageEnhance
from docx import Document
from docx.shared import Inches, Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).parent
ASSET = ROOT / 'docs/inzan_manual_assets'
OUT = ROOT / 'docs/INZAN_ATHLETICS_USER_GUIDE_OPERATIONAL_MANUAL.docx'
SOURCE = ROOT / 'docs/INZAN_ATHLETICS_USER_GUIDE_OPERATIONAL_MANUAL_SOURCE.md'
FONTS = Path('C:/Windows/Fonts')

def font(n, weight='Regular', body=False):
    return ImageFont.truetype(str(FONTS / f'{"OpenSans" if body else "Montserrat"}-{weight}.ttf'), n)

def spaced(draw, xy, text, f, fill, gap=5):
    x, y = xy
    for ch in text:
        draw.text((x,y), ch, font=f, fill=fill)
        x += draw.textlength(ch,font=f)+gap

def cover():
    w,h=1654,2339
    im=Image.new('RGB',(w,h),'#080808')
    photo=Image.open(ASSET/'website_facility.jpg').convert('RGB')
    photo=ImageOps.fit(photo,(w,1580),centering=(.62,.5))
    photo=ImageOps.grayscale(photo).convert('RGB')
    photo=ImageEnhance.Brightness(photo).enhance(.67)
    im.paste(photo,(0,0))
    overlay=Image.new('RGBA',(w,h),(0,0,0,0))
    od=ImageDraw.Draw(overlay)
    for y in range(h):
        alpha= int(230*max(0,min(1,(y-720)/930))) if y<1650 else 235
        od.line((0,y,w,y),fill=(5,5,5,alpha))
    im=Image.alpha_composite(im.convert('RGBA'),overlay)
    logo=Image.open(ROOT/'public/inzanlogo_white.png').convert('RGBA')
    logo.thumbnail((870,250))
    im.alpha_composite(logo,(143,78))
    d=ImageDraw.Draw(im)
    spaced(d,(151,1040),'INZAN ATHLETICS',font(21,'SemiBold'),'#cccccc',5)
    for y,text in [(1122,'USER GUIDE'),(1241,'& OPERATIONAL'),(1360,'MANUAL')]:
        d.text((143,y),text,font=font(103,'ExtraBold'),fill='white',stroke_width=0)
    d.line((150,1605,1500,1605),fill='#555555',width=2)
    d.text((151,1658),'Member services, staff workflows',font=font(31,body=True),fill='#dddddd')
    d.text((151,1706),'and club administration.',font=font(31,body=True),fill='#dddddd')
    spaced(d,(151,2103),'EDITION 03  /  OCTOBER 2026',font(20,'Medium'),'#bcbcbc',3)
    spaced(d,(151,2160),'GARDEN 8  /  NEW CAIRO',font(18,'Medium'),'#888888',3)
    path=ASSET/'editorial_cover.png'; im.convert('RGB').save(path)
    return path

def text_center(d,xy,text,f,fill='#111111'):
    x,y=xy; d.text((x-d.textlength(text,font=f)/2,y),text,font=f,fill=fill)

def arrow(d,points,color='#666666'):
    d.line(points,fill=color,width=3)
    x,y=points[-1]; px,py=points[-2]
    if y>py: tri=[(x,y),(x-8,y-12),(x+8,y-12)]
    elif y<py: tri=[(x,y),(x-8,y+12),(x+8,y+12)]
    else: tri=[(x,y),(x-12,y-8),(x-12,y+8)]
    d.polygon(tri,fill=color)

def diagram(kind):
    im=Image.new('RGB',(1600,600),'white'); d=ImageDraw.Draw(im)
    head=font(32,'SemiBold'); body=font(27,body=True); small=font(23,body=True)
    if kind=='ecosystem':
        for x,t,s in [(265,'MEMBER','Own services'),(800,'COACH','Assigned work'),(1335,'STAFF','Permitted operations')]:
            text_center(d,(x,28),t,head); text_center(d,(x,82),s,body,'#777777')
            arrow(d,[(x,139),(x,204),(800,204),(800,249)])
        d.rectangle((420,252,1180,367),fill='#1a1a1a')
        text_center(d,(800,285),'ONE MEMBER RECORD',font(35,'Bold'),'white')
        for x,t,s in [(190,'Payments','Money collected'),(595,'Entitlements','Service access'),(1000,'Bookings','Reserved time'),(1410,'Attendance','Recorded outcome')]:
            arrow(d,[(800,369),(800,421),(x,421),(x,468)])
            text_center(d,(x,485),t,head); text_center(d,(x,538),s,small,'#666666')
    elif kind=='payment':
        d.rectangle((450,15,1150,109),fill='#1a1a1a')
        text_center(d,(800,42),'CHECK THE PAYMENT RESULT',head,'white')
        for x,t,b1,b2 in [(265,'PAID','Confirm package access','and receipt'),(800,'PENDING','Resolve the payment','before promising access'),(1335,'FAILED','Check the reason','before retrying')]:
            arrow(d,[(800,112),(800,167),(x,167),(x,215)])
            text_center(d,(x,238),t,head)
            d.line((x-195,302,x+195,302),fill='#cccccc',width=2)
            text_center(d,(x,332),b1,body); text_center(d,(x,376),b2,body)
        text_center(d,(800,502),'A saved payment intent is not confirmed service access.',small,'#777777')
    elif kind=='waitlist':
        for x,t in [(240,'ELIGIBILITY'),(800,'CLASS CAPACITY'),(1360,'BOOKING STATE')]:
            text_center(d,(x,12),t,font(23,'SemiBold'),'#777777')
        d.rectangle((70,142,420,252),fill='#1a1a1a');text_center(d,(245,177),'Eligible member',head,'white')
        arrow(d,[(423,197),(580,197)])
        d.rectangle((586,140,1007,254),outline='#555555',width=2);text_center(d,(797,177),'Space available?',head)
        arrow(d,[(1010,172),(1170,172)])
        text_center(d,(1090,129),'YES',small,'#777777');text_center(d,(1360,153),'Confirmed seat',head)
        arrow(d,[(797,257),(797,355),(1168,355)])
        text_center(d,(718,286),'NO',small,'#777777');text_center(d,(1360,338),'Join waitlist',head)
        d.line((72,464,1520,464),fill='#cccccc',width=2)
        text_center(d,(800,488),'A cancellation can promote the first eligible member before the cutoff.',body,'#555555')
        text_center(d,(800,538),'Confirm the new booking state before attending.',small,'#777777')
    elif kind=='handover':
        for i,(title,sub) in enumerate([('RECORD','Payments and refunds'),('COMPARE','Each payment method'),('EXPLAIN','Actual versus expected'),('HAND OVER','Owner and next action')]):
            x=70+i*402
            d.text((x,45),f'0{i+1}',font=font(72,'Light'),fill='#b5b5b5')
            d.line((x,179,x+295,179),fill='#1a1a1a',width=2)
            d.text((x,218),title,font=font(26,'Bold'),fill='#111111')
            for j,word in enumerate(sub.split(' and ')):
                d.text((x,271+j*42),word,font=font(23,body=True),fill='#666666')
        text_center(d,(800,445),'Keep cash, card terminal, bank transfer and Instapay separate.',body,'#555555')
        text_center(d,(800,500),'Explain the difference before closing the shift.',small,'#777777')
    p=ASSET/f'editorial_{kind}.png';im.save(p);return p

def spacing(run,twips=20):
    e=OxmlElement('w:spacing');e.set(qn('w:val'),str(twips));run._element.get_or_add_rPr().append(e)

def no_borders(style):
    ppr=style._element.get_or_add_pPr()
    for el in list(ppr.findall(qn('w:pBdr'))): ppr.remove(el)

def anchor_picture(p,path,width,height):
    inline=p.add_run().add_picture(str(path),width=width,height=height)._inline
    anchor=OxmlElement('wp:anchor')
    for key,val in {'distT':'0','distB':'0','distL':'0','distR':'0','simplePos':'0','relativeHeight':'0','behindDoc':'1','locked':'0','layoutInCell':'1','allowOverlap':'1'}.items():anchor.set(key,val)
    simple=OxmlElement('wp:simplePos');simple.set('x','0');simple.set('y','0');anchor.append(simple)
    for axis in ['H','V']:
        pos=OxmlElement('wp:position'+axis);pos.set('relativeFrom','page');off=OxmlElement('wp:posOffset');off.text='0';pos.append(off);anchor.append(pos)
    anchor.append(deepcopy(inline.find(qn('wp:extent'))));anchor.append(OxmlElement('wp:wrapNone'))
    for tag in ['wp:docPr','wp:cNvGraphicFramePr','a:graphic']:
        node=inline.find(qn(tag))
        if node is not None:anchor.append(deepcopy(node))
    inline.getparent().replace(inline,anchor)

doc=Document()
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3','List Bullet']:
    st=doc.styles[name];st.font.name='Open Sans' if name in ['Normal','List Bullet'] else 'Montserrat'
    st.font.color.rgb=RGBColor.from_string('1A1A1A');no_borders(st)
    st._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),st.font.name);st._element.rPr.rFonts.set(qn('w:hAnsi'),st.font.name)
    for a in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
        st._element.rPr.rFonts.attrib.pop(qn('w:'+a),None)
normal=doc.styles['Normal'];normal.font.size=Pt(10)
normal.paragraph_format.line_spacing=Pt(15);normal.paragraph_format.space_after=Pt(7)
normal.paragraph_format.widow_control=True
for n,size in [('Title',30),('Heading 1',27),('Heading 2',11)]:
    st=doc.styles[n];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.keep_with_next=True
doc.styles['Heading 2'].paragraph_format.space_before=Pt(14);doc.styles['Heading 2'].paragraph_format.space_after=Pt(6)
doc.styles['List Bullet'].font.size=Pt(10);doc.styles['List Bullet'].paragraph_format.line_spacing=Pt(15)
doc.styles['List Bullet'].paragraph_format.space_after=Pt(5)

sec=doc.sections[0];sec.page_width=Mm(210);sec.page_height=Mm(297)
sec.top_margin=Mm(0);sec.bottom_margin=Mm(0);sec.left_margin=Mm(0);sec.right_margin=Mm(0)
sec.header_distance=Mm(0);sec.footer_distance=Mm(0)
anchor_picture(sec.header.paragraphs[0],cover(),Mm(210),Mm(297))
p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(1);p.add_run(' ').font.size=Pt(1)
sec=doc.add_section(WD_SECTION_START.NEW_PAGE)
sec.header.is_linked_to_previous=False;sec.footer.is_linked_to_previous=False
sec.top_margin=Mm(24);sec.bottom_margin=Mm(20);sec.left_margin=Mm(20);sec.right_margin=Mm(20)
sec.header_distance=Mm(10);sec.footer_distance=Mm(10)
for sn in ('Header','Footer'):
    doc.styles[sn].paragraph_format.tab_stops.clear_all()
logo=Image.open(ROOT/'public/inzanlogo.png').convert('RGBA')
logo.crop(logo.getchannel('A').getbbox()).save(ASSET/'header_wordmark.png')
hp=sec.header.paragraphs[0];hp.paragraph_format.space_after=Pt(0);hp.paragraph_format.line_spacing=1.0
hp.add_run().add_picture(str(ASSET/'header_wordmark.png'),width=Mm(47))
hp.paragraph_format.tab_stops.add_tab_stop(Mm(170),WD_ALIGN_PARAGRAPH.RIGHT)
r=hp.add_run('\tOPERATIONS MANUAL');r.font.name='Montserrat';r.font.size=Pt(7);r.font.color.rgb=RGBColor.from_string('777777');spacing(r,20)
fp=sec.footer.paragraphs[0];fp.paragraph_format.tab_stops.add_tab_stop(Mm(170),WD_ALIGN_PARAGRAPH.RIGHT)
r=fp.add_run('EDITION 03   /   OCTOBER 2026\t');r.font.name='Montserrat';r.font.size=Pt(7);r.font.color.rgb=RGBColor.from_string('777777');spacing(r,15)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');fp._p.append(f)

def kicker(text):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(10)
    r=p.add_run(text.upper());r.font.name='Montserrat';r.font.size=Pt(8);r.bold=True;r.font.color.rgb=RGBColor.from_string('777777');spacing(r,32)

def page_title(label,title,deck=None):
    kicker(label)
    p=doc.add_paragraph(title.upper(),style='Heading 1');p.paragraph_format.space_after=Pt(15);p.paragraph_format.line_spacing=1.02
    if deck:
        p=doc.add_paragraph(deck);p.paragraph_format.space_after=Pt(14)
        for r in p.runs:r.font.size=Pt(11);r.font.color.rgb=RGBColor.from_string('555555')

def image(kind):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(12);p.paragraph_format.keep_with_next=False;p.paragraph_format.line_spacing=1.0
    shape=p.add_run().add_picture(str(diagram(kind)),width=Mm(170))
    shape._inline.docPr.set('descr',{'ecosystem':'Member, coach and staff access one shared member record. Payments, entitlements, bookings and attendance remain linked.','payment':'Paid leads to entitlement and receipt. Pending or failed payments require resolution before service access.','waitlist':'Eligible members book available seats; a full class offers a waitlist that may promote before the cutoff.','handover':'Record, compare, explain and hand over each payment method.'}[kind])

def line(text):
    if text.startswith('[[FIG:') or not text.strip():return
    if text.startswith('### '):
        heading=re.sub(r'^\d+\.\d+\s+','',text[4:])
        doc.add_paragraph(heading,style='Heading 2');return
    if text.startswith('- [ ] '):
        p=doc.add_paragraph('☐  '+text[6:]);p.paragraph_format.space_after=Pt(5);return
    if text.startswith('- '):
        p=doc.add_paragraph(style='List Bullet');content=text[2:]
        if ':' in content and len(content.split(':')[0])<50:
            a,b=content.split(':',1);p.add_run(a+':').bold=True;p.add_run(b)
        else:p.add_run(content)
        return
    if re.match(r'^\d+\. ',text):
        n,content=text.split('. ',1);p=doc.add_paragraph();p.paragraph_format.left_indent=Mm(8);p.paragraph_format.first_line_indent=Mm(-8)
        r=p.add_run(n.zfill(2)+'  ');r.bold=True;r.font.name='Montserrat';p.add_run(content);return
    doc.add_paragraph(text)

sections={};key=None
for raw in SOURCE.read_text(encoding='utf-8').splitlines():
    if re.match(r'^## \d+ ',raw):key=int(raw.split()[1]);sections[key]=[]
    elif raw.startswith('## '):key=None
    elif key is not None:sections[key].append(raw)

def blocks(n,start=None,end=None):
    rows=sections[n];out=[];active=start is None
    for row in rows:
        if row.startswith('### '):
            sub=re.match(r'^### (\d+\.\d+)',row)
            tag=sub.group(1) if sub else ''
            if tag==start:active=True
            if tag==end:break
        if active:out.append(row)
    return out

def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,w in enumerate(widths):t.columns[i].width=Mm(w)
    pr=t._tbl.tblPr;bs=OxmlElement('w:tblBorders')
    for edge in ['top','bottom','left','right','insideH','insideV']:
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');bs.append(el)
    pr.append(bs)
    for ri,vals in enumerate([headers]+rows):
        cells=t.rows[0].cells if ri==0 else t.add_row().cells
        for ci,val in enumerate(vals):
            c=cells[ci];c.width=Mm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'110');el.set(qn('w:type'),'dxa');m.append(el)
            tc.append(m);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'1A1A1A' if ri==0 else ('F3F3F3' if ri%2==0 else 'FFFFFF'));tc.append(shade)
            p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(13)
            r=p.add_run(val);r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '1A1A1A');r.bold=ri==0
        props=t.rows[ri]._tr.get_or_add_trPr();props.append(OxmlElement('w:cantSplit'))
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

# Every item below starts on a deliberate page; text remains editable in Word.
plan=[
('01 / THE PLATFORM','Three workspaces\nOne member record',1,None,'1.2','ecosystem','Choose the workspace for your role. The member record connects the work.'),
('01 / ACCESS','Roles and access',1,'1.2',None,None,'Use your own account. Confirm the selected member and branch before acting.'),
('02 / DAILY OPERATIONS','The working day',2,None,None,None,'Open with the schedule. Close with reconciled collections and a clear handover.'),
('03 / MEMBERS','From first enquiry\nto active member',3,None,None,None,'Keep the lead, the sale and the service history on one profile.'),
('04 / SALES','Products and\ncheckout',4,None,'4.4',None,'The final check is the receipt, package and service access—not the submit button.'),
('04 / PAYMENT EXCEPTIONS','Read the result',4,'4.4',None,'payment',None),
('05 / MEMBER ESSENTIALS','The member app',5,None,'5.5',None,'A practical guide to the pass, memberships, classes and PT bookings.'),
('05 / MEMBER SERVICES','Beyond the booking',5,'5.5',None,None,'Nutrition, points, progress and club services live alongside the training schedule.'),
('06 / COACHING','The coach workspace',6,None,None,None,'Keep availability, attendance and follow-up current throughout the day.'),
('07 / BOOKING RULES','A seat is confirmed\nwhen the status says so',7,None,'7.4','waitlist',None),
('07 / ATTENDANCE','Session credits\nand calendar sync',7,'7.4',None,None,'Use the session outcome to explain a balance. Use the booking record to resolve a dispute.'),
('08 / FRONT DESK','Arrival and\nclub operations',8,None,None,None,'Check-in, orders, lockers, guests and member issues share the same working day.'),
('09 / NUTRITION','The consultation\nworkflow',9,None,None,None,'Manage the appointment first. Keep the consultation record in the nutrition workspace.'),
('10 / FINANCE','Collections and\nthe daily close',10,None,None,'handover',None),
('11 / FOLLOW-UP','Calls, tasks\nand notifications',11,None,None,None,'Every open item needs an owner and a next action.'),
('12 / MANAGEMENT','Read the right\nreport',12,None,None,None,'Choose the period and scope before comparing a total.'),
('13 / ADMINISTRATION','Configure the club',13,None,'13.5',None,'Roles, branches, products and member-facing settings.'),
('13 / DATA CONTROL','Backups and imports',13,'13.5',None,None,'Review the file, the tenant and the records that will change before importing.'),
('14 / FACILITIES','Stock and\nequipment support',14,None,None,None,'Service availability and the operational records to maintain.'),
('15 / SECURITY','Accounts, privacy\nand offline work',15,None,None,None,'Confirm the saved result before repeating an action after a connection loss.'),
('16 / EXCEPTIONS','Resolve the cause',16,None,None,None,'Start with the affected record and its status. Keep the reference for escalation.'),
('17 / WORKING CHECKLISTS','Before you finish',17,None,None,None,'Use these checks at the point of work.'),
('18 / RELEASE NOTES','Availability\nand terminology',18,None,None,None,None),
]

page_title('IN THIS HANDBOOK','Find your workflow', 'A desk reference for the people running Inzan Athletics.')
entries=[('01','Workspaces and access',3),('02','Daily operations',5),('03','Leads, members and contracts',6),('04','Products, payments and approvals',7),('05','Member app and services',9),('06','Coach workspace',11),('07','Bookings, attendance and calendar',12),('08','Front desk and club operations',14),('09','Nutrition',15),('10','Finance and payouts',16),('11','Sales follow-up and tasks',17),('12','Dashboards and reporting',18),('13','Administration and data control',19),('14','Inventory and equipment',21),('15','Security and offline work',22),('16','Troubleshooting',23),('17','Working checklists',24),('18','Availability and glossary',25)]
for num,title,pageno in entries:
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(8);p.paragraph_format.tab_stops.add_tab_stop(Mm(170),WD_ALIGN_PARAGRAPH.RIGHT)
    r=p.add_run(num+'   ');r.font.name='Montserrat';r.bold=True;r.font.color.rgb=RGBColor.from_string('888888')
    p.add_run(title+'\t');r=p.add_run(str(pageno).zfill(2));r.bold=True;r.font.name='Montserrat'
doc.add_paragraph('Menus depend on role and enabled features. Where a screen or integration is not yet available, the relevant chapter says so. Use the club’s approved policy for any business decision that remains unresolved.').paragraph_format.space_before=Pt(15)

for label,title,n,start,end,fig,deck in plan:
    doc.add_page_break();page_title(label,title,deck)
    if fig:image(fig)
    for row in blocks(n,start,end):line(row)
    if n==7 and start=='7.4':
        table(['SESSION OUTCOME','CREDIT EFFECT'],[
            ['Attended / completed','Uses one session'],['No Show','Uses one session'],['Rescheduled','No session used'],['Cancellation with at least 12 hours of notice','Credit preserved'],['Cancellation with less than 12 hours of notice','One session forfeited']], [108,62])
        doc.add_paragraph('After a correction',style='Heading 2')
        doc.add_paragraph('Reopen the member’s package and session history. Check the new status, remaining sessions and recorded reason together. Escalate a mismatch with the session ID rather than changing the balance by hand.')
    if n==9:
        doc.add_paragraph('Before saving a consultation',style='Heading 2')
        for s in ['Confirm the member and appointment date.','Check the units and measurements you entered.','Review whether saving will also mark the appointment Completed.','Assign the follow-up action and protect private notes.']:line('- '+s)
    if n==12:
        table(['MEASURE','CHECK BEFORE USING IT'],[['Revenue','Money collected, status exclusions, method and period'],['Attendance','Attended separately from booked and cancelled'],['Class occupancy','Booked seats against the configured capacity'],['Coach earnings','Draft, approved or paid payout state'],['Sales conversion','The documented limitation in the lead denominator']],[52,118])
    if n==13 and start=='13.5':
        doc.add_paragraph('Import procedure',style='Heading 2')
        for i,s in enumerate(['Confirm the source file belongs to Inzan and make a current export before a restore.','Upload the file and review the column mapping. Check name, phone, IDs, dates and amounts.','Compare the preview count with the source. Resolve obvious duplicates before starting the import.','Run the import once. Review its result and Import History before attempting it again.','Open several affected members and verify payments, packages and attendance. Retain the job reference for any exception.'],1):line(f'{i}. {s}')
        doc.add_paragraph('Backup Station recovery',style='Heading 2')
        doc.add_paragraph('Compare offline receipts and attendance entries with the live records before merging the station’s JSON file. A payment recorded during an outage may already have reached the server. Reconcile the amount, member and timestamp before importing it again.')
    if n==14:
        table(['AREA','CURRENT OPERATOR ROUTE'],[['Inventory service','Approved stock register until a dedicated screen is released'],['Equipment service','Approved facility register; report repair and service needs'],['Recurring checklists','Visible Tasks queue; confirm template automation with an administrator']],[52,118])
    if n==18:
        doc.add_paragraph('Terms used in this handbook',style='Heading 2')
        for row in blocks(19):line(row)
        doc.add_paragraph('Document reference',style='Heading 2')
        p=doc.add_paragraph('Edition 03 · October 2026. Workflow content is based on the Inzan CRM repository and its release notes. Visual direction and facility photograph: inzanathletics.com. Prices, thresholds and payout terms remain subject to Inzan management approval.')
        for r in p.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('777777')

doc.core_properties.title='INZAN ATHLETICS — User Guide & Operational Manual'
doc.core_properties.author='Inzan Athletics'
doc.core_properties.subject='Member services and club operations'
doc.core_properties.comments='Edition 03. Website-led editorial design; black, charcoal and white.'
doc.save(OUT)
print(f'{OUT}\nPlanned pages: {len(plan)+2}\nParagraphs: {len(doc.paragraphs)}')
