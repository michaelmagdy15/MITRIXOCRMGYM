from pathlib import Path
import re
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent
SOURCE = BASE / 'docs/INZAN_ATHLETICS_USER_GUIDE_OPERATIONAL_MANUAL_SOURCE.md'
OUTPUT = BASE / 'docs/INZAN_ATHLETICS_USER_GUIDE_OPERATIONAL_MANUAL.docx'
ASSETS = BASE / 'docs/inzan_manual_assets'
ASSETS.mkdir(exist_ok=True)

BLACK = (0, 0, 0)
CHARCOAL = (26, 26, 26)
RAISED = (43, 43, 43)
WHITE = (255, 255, 255)
PALE = (238, 238, 238)
MID = (125, 125, 125)
FONT = 'C:/Windows/Fonts/arial.ttf'
FONT_BOLD = 'C:/Windows/Fonts/arialbd.ttf'

FIGURES = {
    'architecture': ('One member, one operating record', [
        ('Member app', 'Bookings, pass, wallet'), ('Coach portal', 'Schedule and attendance'),
        ('Staff CRM', 'Sales, finance, service'), ('Rules and services', 'Role and transaction checks'),
        ('Inzan data', 'Isolated Firestore')]),
    'daily': ('The daily club operating loop', [
        ('Open', 'Schedule and balances'), ('Serve', 'Check-in and sales'),
        ('Operate', 'Bookings and requests'), ('Close', 'Reconcile and report'),
        ('Handover', 'Owners and next steps')]),
    'lifecycle': ('Lead to active member', [
        ('Capture', 'Lead and source'), ('Qualify', 'Calls and follow-up'),
        ('Convert', 'Single member record'), ('Sell', 'Payment and package'),
        ('Serve', 'Bookings and retention')]),
    'payment': ('Payment and access decision', [
        ('Select', 'Member and product'), ('Validate', 'Identity and price'),
        ('Record', 'Payment transaction'), ('Resolve', 'Paid or pending'),
        ('Confirm', 'Entitlement and receipt')]),
    'booking': ('Booking and attendance lifecycle', [
        ('Explore', 'Time and service'), ('Check', 'Credit and capacity'),
        ('Book', 'Confirmed or waitlist'), ('Attend', 'Check-in or outcome'),
        ('Reconcile', 'Credit, audit, notice')]),
    'reconciliation': ('Shift reconciliation by method', [
        ('Collect', 'Payments and refunds'), ('Group', 'Cash, card, transfer'),
        ('Compare', 'Expected vs actual'), ('Explain', 'Variance and reason'),
        ('Export', 'CSV and handover')]),
}

def make_figure(key):
    title, items = FIGURES[key]
    im = Image.new('RGB', (1500, 360), WHITE)
    draw = ImageDraw.Draw(im)
    title_font = ImageFont.truetype(FONT_BOLD, 37)
    head_font = ImageFont.truetype(FONT_BOLD, 24)
    body_font = ImageFont.truetype(FONT, 20)
    draw.text((50, 25), title, font=title_font, fill=BLACK)
    for i, (head, sub) in enumerate(items):
        x = 44 + i * 289
        fill = CHARCOAL if i % 2 == 0 else RAISED
        draw.rounded_rectangle((x, 137, x + 245, 273), radius=18, fill=fill)
        hb = draw.textbbox((0, 0), head, font=head_font)
        sb = draw.textbbox((0, 0), sub, font=body_font)
        draw.text((x + (245 - (hb[2] - hb[0])) // 2, 173), head, font=head_font, fill=WHITE)
        draw.text((x + (245 - (sb[2] - sb[0])) // 2, 211), sub, font=body_font, fill=WHITE)
        if i < 4:
            draw.line((x + 249, 204, x + 280, 204), fill=MID, width=7)
            draw.polygon([(x + 280, 204), (x + 268, 194), (x + 268, 214)], fill=MID)
    path = ASSETS / f'{key}_mono.png'
    im.save(path)
    return path

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(.72)
section.bottom_margin = Inches(.66)
section.left_margin = Inches(.8)
section.right_margin = Inches(.8)

for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3'):
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), 'Arial')
    style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')

normal = doc.styles['Normal']
normal.font.size = Pt(10)
normal.paragraph_format.line_spacing = 1.16
normal.paragraph_format.space_after = Pt(7)
doc.styles['Title'].font.size = Pt(27)
doc.styles['Title'].font.bold = True
title_style_ppr = doc.styles['Title']._element.get_or_add_pPr()
for border in title_style_ppr.findall(qn('w:pBdr')):
    title_style_ppr.remove(border)
doc.styles['Heading 1'].font.size = Pt(17)
doc.styles['Heading 1'].font.bold = True
doc.styles['Heading 1'].paragraph_format.space_before = Pt(15)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(8)
doc.styles['Heading 1'].paragraph_format.keep_with_next = True
doc.styles['Heading 2'].font.size = Pt(12.5)
doc.styles['Heading 2'].font.bold = True
doc.styles['Heading 2'].paragraph_format.space_before = Pt(11)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(5)
doc.styles['Heading 2'].paragraph_format.keep_with_next = True

logo_path = BASE / 'public/inzanlogo.png'
cover_logo = doc.add_paragraph()
cover_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
cover_logo.paragraph_format.space_after = Pt(18)
cover_logo.add_run().add_picture(str(logo_path), width=Inches(5.5))

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run('INZAN ATHLETICS')
title_ppr = title._p.get_or_add_pPr()
for border in title_ppr.findall(qn('w:pBdr')):
    title_ppr.remove(border)
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('User Guide and Operational Manual')
run.bold = True
run.font.size = Pt(18)
run.font.color.rgb = RGBColor(43, 43, 43)
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run('Version 2.0  |  7 October 2026  |  Inzan Athletics Operations')
meta.runs[0].font.color.rgb = RGBColor(80, 80, 80)
doc.add_paragraph('A practical reference for members, coaches, front desk, sales, managers and administrators. Covers every visible workspace, the decisions each workflow makes and the checks to complete before handing work to the next person.')
doc.add_paragraph('Document map', style='Heading 1')
for entry in ('System map and daily rhythm', 'Lead, member and payment lifecycle', 'Member and coach portals',
              'Booking, attendance and nutrition', 'Front desk, club operations and finance',
              'Sales, tasks, reporting and administration', 'Security, exceptions, checklists and glossary'):
    doc.add_paragraph(entry, style='List Bullet')
doc.add_page_break()

def add_figure(key):
    image = make_figure(key)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(image), width=Inches(6.7))
    caption = doc.add_paragraph(f'Figure  {FIGURES[key][0]}')
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(10)
    caption.runs[0].italic = True
    caption.runs[0].font.size = Pt(8.5)
    caption.runs[0].font.color.rgb = RGBColor(80, 80, 80)

lines = SOURCE.read_text(encoding='utf-8').splitlines()
started = False
for raw in lines:
    line = raw.strip()
    if not line:
        continue
    if line == '## 1 System map and navigation':
        started = True
    if not started:
        continue
    if line.startswith('[[FIG:'):
        add_figure(line[6:-2])
    elif line.startswith('## '):
        doc.add_paragraph(line[3:], style='Heading 1')
    elif line.startswith('### '):
        doc.add_paragraph(line[4:], style='Heading 2')
    elif re.match(r'^\d+\. ', line):
        p = doc.add_paragraph(line)
        p.paragraph_format.left_indent = Inches(.22)
        p.paragraph_format.first_line_indent = Inches(-.22)
    elif line.startswith('- [ ] '):
        doc.add_paragraph('☐ ' + line[6:], style='List Bullet')
    elif line.startswith('- '):
        doc.add_paragraph(line[2:], style='List Bullet')
    else:
        p = doc.add_paragraph(line)
        p.paragraph_format.widow_control = True

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('INZAN ATHLETICS  |  User Guide and Operational Manual  |  ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
for run in footer.runs:
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(80, 80, 80)

doc.core_properties.title = 'INZAN ATHLETICS User Guide and Operational Manual'
doc.core_properties.subject = 'Member, coach, CRM and operations reference'
doc.core_properties.author = 'Inzan Athletics Operations'
doc.save(OUTPUT)
print(f'Created {OUTPUT}; {len(doc.paragraphs)} paragraphs, {len(doc.inline_shapes)} images')
